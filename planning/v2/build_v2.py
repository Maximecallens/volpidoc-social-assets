import sys, datetime as dt
sys.path.insert(0, __import__("os").path.dirname(__file__))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from dev_notes import DEV
from new_rows import REDATE, NEW

SRC = __import__("os").path.join(__import__("os").path.dirname(__file__), "..", "build_planning.py")
src = open(SRC).read()
ns = {}; exec(src.split("wb = Workbook()")[0], ns)
OLD = {r[0]: r for r in ns["rows"]}
RAW = ns["RAW"]

F = "Arial"; NAVY, BLUE, ORANGE = "003046", "0094B6", "F15A24"
thin = Side(style="thin", color="D5E3E5"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
JOURS = ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche"]

def nbsp(t):
    for ch in "?:!;":
        t = t.replace(" " + ch, " " + ch)
    return t

def fr(d):
    x = dt.date.fromisoformat(d); return f"{JOURS[x.weekday()]} {x.day:02d}/{x.month:02d}"

lines = []  # (N°, date, jour, dépôt, fichier, série, message, s1..s4, anim, durée, insp, yt, statut, remarques)
for n in range(1, 12):
    _, d0, dep0, serie, msg, s1, s2, s3, s4, anim, insp, yt, st = OLD[n]
    d, dep = REDATE.get(n, (d0, dep0))
    if n == 1:
        fichier = "short-01 (déjà produit)"
    elif n <= 9:  # déjà livrées : le nom du fichier livré ne change pas
        fichier = f"VD_SH_{n:02d}_{d0}.mp4"
    else:
        fichier = f"VD_SH_{n:02d}_{d}.mp4"
    rem = ""
    if n in DEV:
        st, rem = DEV[n]
        if d != d0:
            st = st + f" · publication avancée au {fr(d)} (nom de fichier inchangé)"
    if n in (10, 11):
        s1, s2, s3, s4, msg, yt = map(nbsp, (s1, s2, s3, s4, msg, yt))
        if n == 10:
            st = "À produire"; rem = ""
    lines.append([f"VD_SH_{n:02d}", d, fr(d).split()[0], dep, fichier, serie, msg, s1, s2, s3, s4, anim, "20 à 24 s", insp, yt, st, rem])
for (n, d, dep, serie, msg, s1, s2, s3, s4, anim, insp, yt) in NEW:
    lines.append([f"VD_SH_{n:02d}", d, fr(d).split()[0], dep, f"VD_SH_{n:02d}_{d}.mp4", serie, nbsp(msg), nbsp(s1), nbsp(s2), nbsp(s3), nbsp(s4), anim, "20 à 24 s", insp, nbsp(yt), "À produire", ""])
lines.sort(key=lambda l: l[1])

wb = Workbook(); ws = wb.active; ws.title = "Planning"
hdr = ["N°", "Date de publication", "Jour", "Dépôt au plus tard (vendredi 18 h)", "Nom du fichier attendu", "Série", "Message clé",
       "Scène 1 — accroche (0–4 s)", "Scène 2 — démonstration (4–11 s)", "Scène 3 — conseil concret (11–17 s)", "Scène 4 — Volpi + appel à l'action (17–22 s)",
       "Animations", "Durée cible", "Inspiration (vidéo déjà déposée)", "Titre YouTube (rédigé par Claude)", "Statut", "Remarques du développeur"]
widths = [10, 13, 10, 14, 30, 20, 34, 40, 52, 44, 48, 38, 10, 26, 40, 22, 40]
ws.append(hdr)
for l in lines: ws.append(l)
for c, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(c)].width = w
for cell in ws[1]:
    cell.font = Font(name=F, bold=True, color="FFFFFF", size=10); cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center"); cell.border = B
ws.row_dimensions[1].height = 42
edit = PatternFill("solid", fgColor="FFF4CC"); done = PatternFill("solid", fgColor="EEF3F4"); wed = PatternFill("solid", fgColor="EAF6F5")
for r in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in r:
        cell.font = Font(name=F, size=10, color=NAVY); cell.alignment = Alignment(wrap_text=True, vertical="top"); cell.border = B
    r[0].font = Font(name=F, size=10, bold=True, color=BLUE)
    r[4].font = Font(name="Courier New", size=10, bold=True, color=ORANGE)
    if r[2].value == "mercredi":
        r[2].fill = wed; r[1].fill = wed
    r[15].fill = edit; r[16].fill = edit
    ws.row_dimensions[r[0].row].height = 150
    if r[15].value and (str(r[15].value).startswith("Déposée") or r[15].value.startswith("Fait par")):
        for c in range(0, 15): r[c].fill = done
        r[0].font = Font(name=F, size=10, bold=True, color="7A8C94")
ws.freeze_panes = "F2"
ws.auto_filter.ref = f"A1:Q{ws.max_row}"
dv = DataValidation(type="list", formula1='"À produire,En cours,Livré,À corriger,Fait par Claude (exemple de référence)"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"P2:P{ws.max_row}")
leg = ws.max_row + 2
notes = [
 ("Légende", None),
 ("Rythme depuis le 26/10/2026 : 2 vidéos par semaine, le lundi et le mercredi (lignes du mercredi sur fond turquoise pâle). Chaque vidéo part sur YouTube Shorts, en Reel Facebook et Instagram, et en vidéo LinkedIn.", 30),
 ("Lignes grisées : vidéos déjà livrées et programmées. Leur date de publication a été avancée pour tenir le nouveau rythme ; le nom de fichier livré ne change pas, rien n'est à refaire.", 30),
 ("Cases jaunes (Statut, Remarques) : à remplir par le développeur. Le reste est rédigé par Volpidoc. La ligne VD_SH_01 sert d'exemple : c'est le Short déjà produit (voir l'onglet « Charte et specs » pour la vidéo de référence).", 30),
 ("Fêtes de fin d'année : les vidéos du 21/12 au 06/01 sont attendues au plus tard le vendredi 18/12. Livrer en avance est toujours possible (idéalement 4 vidéos par vendredi).", 30),
]
for i, (t, h) in enumerate(notes):
    c = ws.cell(row=leg + i, column=1, value=t)
    if h is None:
        c.font = Font(name=F, bold=True, color=NAVY)
    else:
        c.font = Font(name=F, size=10, color=NAVY); c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=leg + i, start_column=1, end_row=leg + i, end_column=11); ws.row_dimensions[leg + i].height = h

# ---- Charte et specs : reprise de l'onglet existant, avec le nouveau rythme
spec = None
exec_ns = {"RAW": RAW, "NAVY": NAVY, "BLUE": BLUE, "ORANGE": ORANGE}
part = src.split("spec = [")[1].split("\n]\n")[0]
spec = eval("[" + part + "\n]", exec_ns)
upd = {
 "Nom du fichier": "Exactement celui de la colonne « Nom du fichier attendu » (ex. VD_SH_12_2026-11-30.mp4 : numéro, puis date de publication, lundi ou mercredi). Une version corrigée : même nom + _v2, _v3… (ex. VD_SH_12_2026-11-30_v2.mp4).",
 "Quand": "Au plus tard le vendredi 18 h indiqué dans la colonne « Dépôt au plus tard » : les vidéos du lundi ET du mercredi de la semaine suivante. Les vidéos peuvent être livrées en avance.",
 "Ce qui se passe ensuite": "Chaque samedi matin, Claude récupère les vidéos déposées, vérifie le format, la charte et les textes, puis programme leur publication le lundi ou le mercredi indiqué (YouTube Shorts, Reels Facebook et Instagram, vidéo LinkedIn). Si une correction est nécessaire, Maxime vous transmet une consigne précise ; la version _v2 est reprise le samedi suivant ou dès qu'elle est déposée avant le samedi.",
}
spec = [(k, upd.get(k, v)) for k, v in spec]
i = [k for k, _ in spec].index("Où")
spec.insert(i, ("Rythme", "2 vidéos par semaine : lundi et mercredi (depuis le 26/10/2026). Le planning est tenu au moins 6 semaines à l'avance ; de nouvelles lignes sont ajoutées au fil de l'eau."))
cs = wb.create_sheet("Charte et specs")
cs.column_dimensions["A"].width = 30; cs.column_dimensions["B"].width = 110; cs.column_dimensions["C"].width = 8
sw = {"Texte principal": NAVY, "Accent (mots clés)": BLUE, "Orange (dates, alertes, chiffres clés)": ORANGE, "Turquoise (décor, cases, puces)": "26A7BA"}
for r, (k, v) in enumerate(spec, start=1):
    a = cs.cell(row=r, column=1, value=k)
    if v is None:
        a.font = Font(name=F, bold=True, color="FFFFFF", size=11); a.fill = PatternFill("solid", fgColor=NAVY)
        cs.cell(row=r, column=2).fill = PatternFill("solid", fgColor=NAVY)
    else:
        b = cs.cell(row=r, column=2, value=v)
        a.font = Font(name=F, bold=True, color=NAVY, size=10); b.font = Font(name=F, color=NAVY, size=10)
        a.alignment = b.alignment = Alignment(wrap_text=True, vertical="top"); a.border = b.border = B
        cs.row_dimensions[r].height = max(18, 15 * (len(v) // 100 + 1))
    if k in sw: cs.cell(row=r, column=3).fill = PatternFill("solid", fgColor=sw[k])

out = sys.argv[1]
wb.save(out)
for l in lines: print(l[0], l[1], l[2], l[3], l[4], "|", l[15][:40])
