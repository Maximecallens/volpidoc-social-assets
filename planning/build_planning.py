from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

RAW = "https://raw.githubusercontent.com/Maximecallens/volpidoc-social-assets/main/"
CTA = "Bloc final (fond dégradé bleu, texte blanc) : « Combien vos angles morts vous coûtent‑ils ? » + « volpidoc.com ». Tenue 3 s minimum."
F = "Arial"
NAVY, BLUE, ORANGE = "003046", "0094B6", "F15A24"
thin = Side(style="thin", color="D5E3E5")
B = Border(left=thin, right=thin, top=thin, bottom=thin)

def fname(n, d, v=""): return f"VD_SH_{n:02d}_{d}{v}.mp4"

rows = [
 # n, date, dépôt, série, message, s1, s2, s3, s4, anim, inspiration, titre YT, statut
 (1,"2026-10-19","—","Avant le 31/12",
  "Un contrat renouvelé au 1er janvier se résilie souvent avant le 31 octobre (préavis de 2 mois).",
  "« Votre contrat se renouvelle le 1er janvier. » puis « Le vrai rendez‑vous ? Le 31 octobre. » (31 octobre en orange)",
  "« Pourquoi ? Le préavis. » Frise qui se trace du 1er janvier vers le 31 octobre ; pastille « 2 mois de préavis (souvent) » ; « Sans courrier à temps, le contrat repart pour un an. »",
  "« À vérifier cette semaine : » 3 cases qui se cochent : « La date d'échéance exacte », « La durée du préavis », « La forme exigée (recommandé, e‑mail…) »",
  "Renard plein + « Volpi lit vos contrats et repère vos préavis pour vous. » + "+CTA,
  "Frise animée, cases qui se cochent une à une, fondus enchaînés entre scènes.",
  "—","Contrat renouvelé au 1er janvier : le vrai délai, c'est le 31 octobre #Shorts","Fait par Claude (exemple de référence)"),
 (2,"2026-10-26","2026-10-23","Avant le 31/12",
  "Il reste quelques jours pour prévenir : voici comment résilier proprement.",
  "« Contrat au 1er janvier, préavis de 2 mois ? » puis « Il vous reste 5 jours. » (5 jours en orange)",
  "Calendrier qui décompte de 5 à 0 jusqu'au 31 octobre, puis 3 étapes qui s'affichent une à une : « 1. Retrouver la date et le préavis » « 2. Rédiger la lettre » « 3. L'envoyer avec une preuve »",
  "« La bonne forme ? Celle prévue au contrat : recommandé, e‑mail ou espace client. » puis « Gardez toujours la preuve d'envoi. »",
  "Renard plein + « Volpi retrouve vos préavis et pré‑remplit la lettre de résiliation. » + "+CTA,
  "Compte à rebours sur un calendrier, étapes qui se cochent. Ne pas montrer d'envoi automatique (l'envoi reste à faire par l'utilisateur).",
  "Le préavis était hier","Résilier un contrat avant le 31 octobre : 3 étapes #Shorts","À produire"),
 (3,"2026-11-02","2026-10-30","Le vrai coût de vos contrats",
  "Une clause d'indexation fait grimper le prix chaque année, sans bruit.",
  "« Votre contrat coûtait 1 200 € par an. » puis « Et aujourd'hui ? »",
  "Compteur qui grimpe année par année (+4 % par an) : « Année 1 : 1 248 € » « Année 2 : 1 298 € » « Année 3 : 1 350 € » ; badge orange « +150 € par an (+12,5 %) » ; petite mention « Exemple illustratif ».",
  "« À vérifier : » 3 cases : « La clause d'indexation ou de révision », « Le prix actuel face au prix d'origine », « Les options que vous n'utilisez plus »",
  "Renard plein + « Volpi extrait les clauses clés de vos contrats. » + "+CTA,
  "Compteur de chiffres animé (chiffres à chasse fixe), barres qui montent, badge qui apparaît.",
  "—","Indexation : pourquoi votre contrat coûte plus cher chaque année #Shorts","À produire"),
 (4,"2026-11-09","2026-11-06","Le vrai coût de vos contrats",
  "Deux contrats peuvent couvrir le même risque : vous payez alors deux fois.",
  "« Deux contrats. » « Une garantie commune. » « Payée deux fois. » (3 lignes qui tombent l'une après l'autre)",
  "Deux cartes contrat (« Multirisque pro » et « Protection juridique ») glissent l'une vers l'autre ; la ligne commune « Protection juridique » s'allume en orange dans les deux. Petite mention « Exemple illustratif ».",
  "« Où chercher : » « Des contrats souscrits à des moments différents », « Des options ajoutées par deux interlocuteurs », « Une garantie déjà incluse dans un autre contrat »",
  "Renard plein + « Volpi repère les garanties en doublon entre vos contrats. » + "+CTA,
  "Deux cartes qui se superposent, surlignage orange de la ligne en double.",
  "La ligne payée deux fois","Garanties en double : payez‑vous deux fois le même risque ? #Shorts","À produire"),
 (5,"2026-11-16","2026-11-13","Assurance : lire entre les lignes",
  "Une garantie n'a de valeur qu'avec sa franchise et son plafond.",
  "« Sinistre : 8 000 €. » puis « Vous êtes assuré. Vous touchez combien ? »",
  "Barre de 8 000 € qui se découpe : franchise 1 500 € (orange), indemnité plafonnée à 5 000 € (bleu), reste à charge. Résultat : « Indemnisé : 5 000 € · À votre charge : 3 000 € ». Petite mention « Exemple illustratif ».",
  "« À vérifier : » 3 cases : « Le montant de la franchise », « Le plafond par sinistre et par an », « Les sous‑limites de certaines garanties »",
  "Renard plein + « Volpi résume chaque contrat et ses clauses clés en moins de 30 secondes. » + "+CTA,
  "Barre qui se découpe en segments colorés, chiffres qui apparaissent sur chaque segment.",
  "—","Franchise et plafond : ce que votre assurance paie vraiment #Shorts","À produire"),
 (6,"2026-11-23","2026-11-20","Assurance : lire entre les lignes",
  "Deux devis au même prix peuvent couvrir très différemment.",
  "« Deux devis. » puis « Même prix. » puis « Pas la même couverture. »",
  "Deux colonnes « Devis A » / « Devis B », même prix affiché « 1 800 €/an ». Lignes qui apparaissent : Franchise 500 € / 1 500 € ; Plafond 300 000 € / 150 000 € ; Protection juridique ✓ / ✗. Les écarts s'allument en orange côté B. Petite mention « Exemple illustratif ».",
  "« Comparez ligne à ligne : franchises, plafonds, exclusions, options. »",
  "Renard plein + « Volpi compare jusqu'à 5 devis que vous importez face à votre contrat actuel. » + "+CTA,
  "Tableau comparatif qui se remplit ligne par ligne, surlignages orange. Ne jamais suggérer que Volpidoc cherche des offres sur le marché.",
  "Même prix, pas la même couverture","Deux devis au même prix : lequel vous couvre vraiment ? #Shorts","À produire"),
 (7,"2026-11-30","2026-11-27","Vos contrats ont la réponse",
  "La réponse à « suis‑je couvert ? » est dans le contrat, encore faut‑il la trouver.",
  "« Suis‑je couvert si… ? » puis « La réponse est page 47. »",
  "Un PDF dont les pages défilent très vite (1 → 47), puis une bulle de chat : question « Suis‑je couvert en cas de dégât des eaux dans mon local ? » ; réponse de Volpi avec une étiquette « Source : multirisque pro, page 47 » (ne pas écrire de réponse juridique précise).",
  "« Une question précise mérite une réponse sourcée. » puis « Vérifiez toujours la page citée. »",
  "Renard plein + « Demandez à Volpi : il répond en citant vos documents. » + "+CTA,
  "Défilement rapide de pages, bulles de chat qui apparaissent, étiquette « Source » qui clignote une fois.",
  "La réponse est page 47 ; Vos contrats répondent","Votre contrat a la réponse… page 47 #Shorts","À produire"),
 (8,"2026-12-07","2026-12-04","Reprendre la main",
  "Avant de piloter ses contrats, il faut savoir combien on en a.",
  "« Combien de contrats a votre entreprise ? » puis « Vraiment ? »",
  "Des tuiles apparaissent une à une : Assurances, Télécom, Énergie, Logiciels, Maintenance, Leasing, Banque… et un compteur qui monte puis affiche « Plus que vous ne le pensez. »",
  "« Commencez par la liste : qui, quoi, combien, jusqu'à quand. »",
  "Renard plein + « Importez vos contrats en PDF, Word ou scan : Volpi les range au même endroit. » + "+CTA,
  "Grille de tuiles qui se remplit, compteur animé.",
  "L'inventaire que personne n'a fait","Combien de contrats a votre entreprise ? Faites l'inventaire #Shorts","À produire"),
 (9,"2026-12-14","2026-12-11","Reprendre la main",
  "Classer ses contrats ne suffit pas : il faut suivre leurs échéances.",
  "« Vos contrats sont tous dans un dossier. » puis « Et vos échéances ? »",
  "Un dossier de PDF s'ouvre et se transforme en frise d'échéances colorées (à surveiller, à renouveler, à résilier).",
  "« Un bon suivi : une date, un préavis, une action. »",
  "Renard plein + « Volpi suit vos échéances et les exporte dans votre agenda. » + "+CTA,
  "Transformation dossier → frise, pastilles de couleur par statut.",
  "Un dossier PDF n'est pas un pilotage ; Quatre contrats, un seul à ouvrir","Classer ses contrats ne suffit pas : suivez vos échéances #Shorts","À produire"),
 (10,"2026-12-21","2026-12-18","Reprendre la main",
  "Une revue de contrats de 20 minutes par trimestre évite la plupart des mauvaises surprises.",
  "« 20 minutes. » puis « Une fois par trimestre. »",
  "Une horloge qui tourne ; 4 cases qui se cochent : « Les échéances des 3 prochains mois », « Les préavis à envoyer », « Les hausses de prix », « Les garanties en double »",
  "« Bloquez le créneau dans votre agenda dès maintenant. »",
  "Renard plein + « Volpi prépare le terrain : échéances, clauses, doublons. » + "+CTA,
  "Horloge animée, cases qui se cochent.",
  "20 minutes par trimestre","Contrats : la revue trimestrielle en 20 minutes #Shorts","À produire"),
 (11,"2026-12-28","2026-12-18","Bilan de fin d'année",
  "Trois réflexes simples pour mieux gérer ses contrats en 2027.",
  "« 2027 : et si vos contrats prenaient de bonnes résolutions ? »",
  "3 cartes qui apparaissent : « 1. Noter chaque date de préavis » « 2. Relire les hausses de prix » « 3. Comparer avant de renouveler »",
  "« Un contrat suivi coûte moins cher qu'un contrat oublié. »",
  "Renard plein + « Volpi lit vos contrats et vous montre ce qui vous échappe. » + « Essai gratuit de 14 jours » + "+CTA,
  "Cartes qui se retournent, petit effet de confettis aux couleurs de la charte (discret).",
  "—","3 réflexes pour mieux gérer vos contrats en 2027 #Shorts","À produire"),
]

wb = Workbook()
ws = wb.active; ws.title = "Planning"
hdr = ["N°","Date de publication (lundi)","Dépôt au plus tard (18 h)","Nom du fichier attendu","Série","Message clé",
       "Scène 1 — accroche (0–4 s)","Scène 2 — démonstration (4–11 s)","Scène 3 — conseil concret (11–17 s)","Scène 4 — Volpi + appel à l'action (17–22 s)",
       "Animations","Durée cible","Inspiration (vidéo déjà déposée)","Titre YouTube (rédigé par Claude)","Statut","Remarques du développeur"]
widths = [10,14,14,30,20,34,40,52,44,48,38,10,26,40,16,30]
ws.append(hdr)
for i,(n,d,dep,serie,msg,s1,s2,s3,s4,anim,insp,yt,st) in enumerate(rows, start=2):
    ws.append([f"VD_SH_{n:02d}", d, dep, fname(n,d) if n>1 else "short-01 (déjà produit)", serie, msg, s1, s2, s3, s4, anim, "20 à 24 s", insp, yt, st, ""])
for c,w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(c)].width = w
for cell in ws[1]:
    cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    cell.border = B
ws.row_dimensions[1].height = 42
edit = PatternFill("solid", fgColor="FFF4CC")
for r in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in r:
        cell.font = Font(name=F, size=10, color=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = B
    r[0].font = Font(name=F, size=10, bold=True, color=BLUE)
    r[3].font = Font(name="Courier New", size=10, bold=True, color=ORANGE)
    r[14].fill = edit; r[15].fill = edit
    ws.row_dimensions[r[0].row].height = 150
ws.cell(row=2, column=1).font = Font(name=F, size=10, bold=True, color="7A8C94")
for c in range(1,15): ws.cell(row=2, column=c).fill = PatternFill("solid", fgColor="EEF3F4")
ws.freeze_panes = "E2"
ws.auto_filter.ref = f"A1:P{ws.max_row}"
dv = DataValidation(type="list", formula1='"À produire,En cours,Livré,À corriger,Fait par Claude (exemple de référence)"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"O2:O{ws.max_row}")
leg = ws.max_row + 2
ws.cell(row=leg, column=1, value="Légende").font = Font(name=F, bold=True, color=NAVY)
ws.cell(row=leg+1, column=1, value="Cases jaunes (Statut, Remarques) : à remplir par le développeur. Le reste est rédigé par Volpidoc. La ligne VD_SH_01 sert d'exemple : c'est le Short déjà produit (voir l'onglet « Charte et specs » pour la vidéo de référence).").font = Font(name=F, size=10, color=NAVY)
ws.merge_cells(start_row=leg+1, start_column=1, end_row=leg+1, end_column=10)
ws.row_dimensions[leg+1].height = 30
ws.cell(row=leg+1, column=1).alignment = Alignment(wrap_text=True, vertical="top")

# ---- Charte et specs
cs = wb.create_sheet("Charte et specs")
cs.column_dimensions["A"].width = 30; cs.column_dimensions["B"].width = 110
spec = [
 ("FORMAT", None),
 ("Fichier", "MP4, vidéo H.264 + audio AAC (piste audio présente, même silencieuse). Moins de 100 Mo."),
 ("Dimensions", "1080 × 1920 px (9:16, vertical), 30 images/s."),
 ("Durée", "20 à 24 s (18 s minimum, 60 s maximum). La scène finale (appel à l'action) reste visible au moins 3 s, sans rien qui bouge à la fin."),
 ("Zone sûre", "Tout texte et élément important entre y = 250 px et y = 1650 px (les interfaces YouTube et Instagram masquent le haut et le bas). Marges latérales : 90 px."),
 ("Son", "Silence ou musique libre de droits uniquement (pas de musique commerciale)."),
 ("CHARTE", None),
 ("Fond", "Clair : dégradé à 165° de #FFFFFF → #EAF6F5 → #D9F0EE. Pas de fond bleu marine dominant."),
 ("Texte principal", "#003046 (bleu marine)"),
 ("Accent (mots clés)", "#0094B6 (bleu Volpidoc)"),
 ("Turquoise (décor, cases, puces)", "#26A7BA et #64C5C1"),
 ("Orange (dates, alertes, chiffres clés)", "#F15A24 — avec parcimonie, un ou deux éléments par scène"),
 ("Police", "Plus Jakarta Sans : 800 pour les titres (80 à 92 px), 700 pour les listes (44 à 50 px), 500 pour les textes secondaires (34 à 46 px). Google Fonts, ou fichiers : "+RAW+"brand/fonts/"),
 ("Logo", "Logo horizontal couleur en haut à gauche (hauteur ≈ 62 px, à y ≈ 260 px) : "+RAW+"brand/logos/logo-horizontal.png — versions : logo-vertical.png, logo-horizontal-bleu-fonce.png, logo-horizontal-typo-blanche.png (sur fond foncé seulement). Ne jamais déformer ni recolorer le logo."),
 ("Renard", "En filigrane (opacité ≈ 9 %) à droite sur toute la vidéo, et en plein (≈ 280 px) en haut de la scène finale : "+RAW+"brand/logos/renard-seul.png"),
 ("En-tête", "Logo à gauche, nom de la série à droite en orange #F15A24, majuscules, 26 px, espacement des lettres 0,14 em."),
 ("Pied de page", "« Volpidoc est une aide à la lecture : ni courtier, ni conseil juridique. » — 26 px, #003046 à 55 % d'opacité, à y ≈ 1600 px, sur toute la vidéo."),
 ("Scène finale", "Renard plein + phrase sur Volpi (voir planning) + bloc arrondi (rayon 32 px) en dégradé #0094B6 → #26A7BA, texte blanc : « Combien vos angles morts vous coûtent‑ils ? » (50 px, 800) et « volpidoc.com » (42 px, 700)."),
 ("Typographie", "Français soigné : espace insécable avant ? : ; ! — guillemets « » — aucun mot coupé en fin de ligne (ex. « rendez‑vous », « e‑mail » sur une seule ligne). Relire l'orthographe."),
 ("RÉFÉRENCE", None),
 ("Vidéo de référence", "Short 01 (conforme) : "+RAW+"posts/short-01_preavis/video.mp4"),
 ("Gabarit réutilisable", "Code HTML de la vidéo de référence (animation image par image, réutilisable tel quel en changeant les textes et les animations) : "+RAW+"posts/short-01_preavis/motion.html — rendu : "+RAW+"tools/motion.py"),
 ("CONTENU", None),
 ("Textes", "Reprendre exactement les textes du planning (entre « »). Toute proposition de reformulation va dans la colonne Remarques."),
 ("Chiffres", "Tout exemple chiffré affiche la mention « Exemple illustratif »."),
 ("Interdits", "Pas de multi‑sociétés / multi‑organisation. Pas de fonctionnalité absente de volpidoc.com. Pas de faux témoignage ni de cas client chiffré. Ne jamais montrer Volpidoc qui cherche des offres sur le marché ou qui envoie un courrier à la place de l'utilisateur. Aucune mention de MedMax, EvoTattoo ou Tattoomorrow."),
 ("LIVRAISON", None),
 ("Nom du fichier", "Exactement celui de la colonne « Nom du fichier attendu » (ex. VD_SH_02_2026-10-26.mp4). Une version corrigée : même nom + _v2, _v3… (ex. VD_SH_02_2026-10-26_v2.mp4)."),
 ("Où", "À la racine de ce drive partagé, à côté de ce planning."),
 ("Quand", "Au plus tard le vendredi 18 h indiqué dans la colonne « Dépôt au plus tard ». Les vidéos peuvent être livrées en avance."),
 ("Ce qui se passe ensuite", "Chaque samedi matin, Claude récupère les vidéos déposées, vérifie le format, la charte et les textes, puis programme la publication du lundi (YouTube Shorts, Reels Facebook et Instagram). Si une correction est nécessaire, Maxime vous transmet une consigne précise ; la version _v2 est reprise le samedi suivant ou dès qu'elle est déposée avant le samedi."),
]
r=1
for k,v in spec:
    a=cs.cell(row=r,column=1,value=k); 
    if v is None:
        a.font=Font(name=F,bold=True,color="FFFFFF",size=11); a.fill=PatternFill("solid",fgColor=NAVY)
        cs.cell(row=r,column=2).fill=PatternFill("solid",fgColor=NAVY)
    else:
        b=cs.cell(row=r,column=2,value=v)
        a.font=Font(name=F,bold=True,color=NAVY,size=10); b.font=Font(name=F,color=NAVY,size=10)
        a.alignment=b.alignment=Alignment(wrap_text=True,vertical="top")
        a.border=b.border=B
        cs.row_dimensions[r].height = max(18, 15*(len(v)//100+1))
    r+=1
# color swatches
sw={"Texte principal":NAVY,"Accent (mots clés)":BLUE,"Orange (dates, alertes, chiffres clés)":ORANGE,"Turquoise (décor, cases, puces)":"26A7BA"}
cs.column_dimensions["C"].width = 8
for row in cs.iter_rows(min_row=1, max_row=cs.max_row):
    if row[0].value in sw: cs.cell(row=row[0].row, column=3).fill = PatternFill("solid", fgColor=sw[row[0].value])

wb.save("/home/claude/vd/planning/Volpidoc - Planning motion design.xlsx")
print("ok", ws.max_row)
